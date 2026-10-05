"""
ONE-TIME: purani products.json  ->  products.xlsx
(GitHub Action 'make-excel.yml' isko chalata hai, ya local: python scripts/json_to_excel.py)
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

JSON_FILE = "products.json"
EXCEL_FILE = "products.xlsx"

HEADERS = ["id", "name", "company", "price", "category", "packing", "saltContent",
           "uses", "bestseller", "inStock", "discount", "scheme", "image"]
WIDTHS = [7, 30, 18, 10, 14, 12, 42, 55, 12, 10, 10, 22, 40]
MAXROW = 2000


def val(p, key):
    v = p.get(key)
    if key == "saltContent" and v is None:
        v = p.get("salt content")          # kuch rows me key space ke saath hai
    return v


def main():
    with open(JSON_FILE, encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(HEADERS)
    cats = []
    for p in data:
        row = []
        for h in HEADERS:
            if h == "image":
                imgs = p.get("images") or ([p["image"]] if p.get("image") else [])
                # folder prefix hata do, sirf file ka naam rakho
                imgs = [i[len("images/"):] if i.startswith("images/") else i for i in imgs]
                row.append(", ".join(imgs) or None)
                continue
            v = val(p, h)
            if isinstance(v, str):
                v = v.strip()
                if h == "price":
                    try:
                        v = float(v)
                        v = int(v) if v.is_integer() else v
                    except ValueError:
                        pass
            row.append(v)
        ws.append(row)
        c = p.get("category")
        if c and c not in cats:
            cats.append(c)

    thin = Side(style="thin", color="BBBBBB")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6F5C")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 300)):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=(c.column in (7, 8, 13)))
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{MAXROW}"

    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    dvc = DataValidation(type="list", formula1='"' + ",".join(cats) + '"', allow_blank=True,
                         errorStyle="warning", showErrorMessage=True,
                         errorTitle="Nayi category?", error="Ye category list me nahi hai. Pakka sahi hai?")
    dvp = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                         showErrorMessage=True, errorTitle="Galat price", error="Price number hona chahiye")
    for d in (dv, dvc, dvp):
        ws.add_data_validation(d)
    dv.add(f"I2:J{MAXROW}")
    dvc.add(f"E2:E{MAXROW}")
    dvp.add(f"D2:D{MAXROW}")

    ins = wb.create_sheet("Instructions")
    lines = [
        ("Excel se Website update - Kaise use karein", True), ("", False),
        ("1. 'Products' sheet me MRP, discount, scheme, stock (inStock) edit karo.", False),
        ("2. Naya product: sabse neeche nayi row me likho. 'id' khali chhodo to automatically agla number mil jayega.", False),
        ("3. bestseller / inStock: dropdown se TRUE ya FALSE chuno. Stock khatam = inStock FALSE.", False),
        ("4. discount / scheme: nahi hai to khali chhodo. Likh sakte ho jaise 10 ya '10+1 free'.", False),
        ("5. image: sirf file ka naam likho (jaise acnetoin_soap.jpg). Ek se zyada photo ho to comma se alag karo: a.jpg, b.jpg", False),
        ("6. Photos GitHub ke 'images' folder me upload karni hain (naam bilkul wahi jo Excel me likha hai).", False),
        ("7. Product hatana ho to puri row delete karo.", False),
        ("8. products.xlsx ko GitHub repo me replace upload karo. 1-2 minute me website update ho jayegi.", False),
        ("", False),
        ("Dhyan: Row 1 ke header ka naam mat badlo. Price sirf number me likho. File ka naam products.xlsx hi rakho.", True),
    ]
    for i, (t, b) in enumerate(lines, start=1):
        ins.cell(row=i, column=1, value=t).font = Font(name="Arial", size=11, bold=b)
    ins.column_dimensions["A"].width = 125
    wb.save(EXCEL_FILE)
    print(f"OK: {len(data)} products -> {EXCEL_FILE}")


if __name__ == "__main__":
    main()
