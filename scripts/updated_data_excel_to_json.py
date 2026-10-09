"""
products.xlsx  ->  products.json   (GitHub Action 'excel-to-json.yml' isko chalata hai)
Local test:  python scripts/updated_data_excel_to_json.py

Naye columns (Excel me kahin bhi, header ka naam 'Description' / 'Rx' kuch bhi case me chalega):
  Description : product ka apna likha description (khali ho to website auto template banati hai)
  Rx          : khali ya Yes = prescription medicine (default). No / OTC / FALSE = Rx nahi.
"""
import json
import os
import sys
from openpyxl import load_workbook

# ====== SETTINGS ======
EXCEL_FILE = "products.xlsx"
SHEET_NAME = "Products"
JSON_FILE = "products.json"
IMAGE_FOLDER = "images"
CHECK_IMAGES = True
# ======================

TEXT_FIELDS = ["name", "company", "category", "packing", "saltContent", "uses"]
NULLABLE_FIELDS = ["discount", "scheme"]
ORDER = ["id", "name", "company", "price", "category", "packing", "saltContent",
         "uses", "bestseller", "inStock", "discount", "scheme", "rx", "description"]

# Header ka naam capital/small kisi bhi tarah likha ho, sahi naam pakad lo
# (jaise "Description" -> "description", "Rx" -> "rx", "instock" -> "inStock")
CANON = {h.lower(): h for h in ORDER + ["image"]}


def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    return v


def to_number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else round(float(v), 2)
    try:
        f = float(str(v).replace(",", "").replace("₹", "").replace("%", "").strip())
        return int(f) if f.is_integer() else round(f, 2)
    except ValueError:
        return None


def to_bool(v, default):
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in ("true", "yes", "y", "1", "haan", "ha"):
        return True
    if s in ("false", "no", "n", "0", "nahi", "na"):
        return False
    return default


def to_rx(v):
    """Rx: khali = Rx (default). Sirf No / OTC / FALSE jaisa kuch likha ho tabhi Rx hat'ta hai."""
    if v is None:
        return True
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    return s not in ("false", "no", "n", "0", "nahi", "na", "otc", "non-rx", "non rx", "nonrx")


def fix_image(name):
    name = name.strip()
    if not name:
        return None
    if "/" in name or name.startswith("http"):
        return name
    return f"{IMAGE_FOLDER}/{name}"


def main():
    if not os.path.exists(EXCEL_FILE):
        sys.exit(f"ERROR: {EXCEL_FILE} nahi mili")
    wb = load_workbook(EXCEL_FILE, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"ERROR: '{SHEET_NAME}' sheet nahi mili. Sheets: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    headers = []
    for c in ws[1]:
        h = str(c.value).strip() if c.value is not None else ""
        headers.append(CANON.get(h.lower(), h))
    for h in ("name", "price"):
        if h not in headers:
            sys.exit(f"ERROR: header me '{h}' column chahiye")

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rec = {h: clean(v) for h, v in zip(headers, r) if h}
        if any(v is not None for v in rec.values()):
            rows.append(rec)

    used = [to_number(r.get("id")) for r in rows]
    next_id = max([u for u in used if u is not None] or [0]) + 1
    products, errors, warnings, seen = [], [], [], set()

    for i, rec in enumerate(rows, start=2):
        name = rec.get("name")
        price = to_number(rec.get("price"))
        if not name:
            errors.append(f"Row {i}: name khali hai"); continue
        if price is None:
            errors.append(f"Row {i} ({name}): price galat/khali hai"); continue

        pid = to_number(rec.get("id"))
        if pid is None:
            pid = next_id; next_id += 1
        if pid in seen:
            errors.append(f"Row {i} ({name}): id {pid} duplicate hai"); continue
        seen.add(pid)

        p = {"id": pid, "price": price}
        for f in TEXT_FIELDS:
            p[f] = rec.get(f) or ""
        p["bestseller"] = to_bool(rec.get("bestseller"), False)
        p["inStock"] = to_bool(rec.get("inStock"), True)
        p["rx"] = to_rx(rec.get("rx"))
        p["description"] = rec.get("description") or ""     # Alt+Enter wali nayi lines bhi bachti hain
        for f in NULLABLE_FIELDS:
            v = rec.get(f)
            if v is not None and not isinstance(v, str):
                v = to_number(v)
            elif isinstance(v, str):
                n = to_number(v)
                v = n if n is not None else v
            p[f] = v

        out = {k: p.get(k) for k in ORDER}

        # image: ek ho to "image", comma se kai ho to "images" list
        imgs = [fix_image(x) for x in str(rec.get("image") or "").split(",")]
        imgs = [x for x in imgs if x]
        if len(imgs) == 1:
            out["image"] = imgs[0]
        elif len(imgs) > 1:
            out["images"] = imgs
        for im in imgs:
            if CHECK_IMAGES and not im.startswith("http") and not os.path.exists(im):
                warnings.append(f"Row {i} ({name}): image nahi mili -> {im}")
        products.append(out)

    for w in warnings:
        print("WARNING:", w)
    if errors:
        for e in errors:
            print("ERROR:", e)
        sys.exit("Excel me galtiyan hain, products.json update NAHI hua.")

    with open(JSON_FILE, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"OK: {len(products)} products -> {JSON_FILE}")


if __name__ == "__main__":
    main()
